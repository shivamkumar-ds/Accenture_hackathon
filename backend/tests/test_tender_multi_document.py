"""
Regression coverage for the multi-document Tender + XLS/XLSX support
feature (real-CPPP-tender validation gap -- a real tender was found to
consist of tender.pdf + tech.xls + BOQ_969057.xls, and BidOps could only
ever attach one PDF to a Tender).

Covers, per the governing spec's Section 8 test list:
  1. PDF-only tender still works (backward compatibility)
  2. XLS upload accepted
  3. XLSX upload accepted
  4. Unsupported file types still rejected cleanly
  5. Multiple documents attached to one Tender
  6. Spreadsheet parsing (single + multiple + empty sheets)
  7. Combined tender extraction (PDF + spreadsheet feeding one pipeline)
  8. Source-document traceability (source_document_id/source_location)
  9. Financial/BOQ-role documents excluded from LLM extraction input
  10. Company isolation (list_tender_documents / add_tender_document)
  11. Existing single-PDF analysis path remains intact

Uses the same seam-mocking pattern established for Bug #005/#006 and the
existing tender tests: an in-memory SQLite database (structural coverage
only -- this project's standing convention, since SQLite lacks some
Postgres-specific DDL semantics, though nothing in this feature depends
on those), and provider="mock" (app/agents/mock_extraction.py) rather
than a real LLM call, driving the *real* tender_analyzer/tender_service
code against real PDF (via reportlab, already a project dependency) and
real .xlsx (via openpyxl) files on disk.
"""

import asyncio
import io
import uuid
from pathlib import Path

import openpyxl
import pytest
from reportlab.pdfgen import canvas
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.agents import document_parser
from app.agents.tender_analyzer import TenderSourceDocument, analyze_tender
from app.core import storage
from app.core.database import Base
from app.models import Company, Document, Mission, Requirement, Tender, User
from app.models.enums import DocumentProcessingStatus, MissionStatus, UserRole, UserStatus
from app.services import document_service, tender_service
from app.services.exceptions import ConflictError, NotFoundError, UnsupportedFileTypeError


# ---------------------------------------------------------------------------
# Shared fixtures
# ---------------------------------------------------------------------------


@pytest.fixture()
def db():
    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(
        engine,
        tables=[
            Company.__table__, User.__table__, Mission.__table__, Tender.__table__,
            Requirement.__table__, Document.__table__,
        ],
    )
    session = sessionmaker(bind=engine)()
    yield session
    session.close()


def _make_company_and_user(db):
    company = Company(id=uuid.uuid4(), name="Acme", registration_number=str(uuid.uuid4()))
    db.add(company)
    db.flush()
    user = User(
        id=uuid.uuid4(), company_id=company.id, name="Admin", email=f"{uuid.uuid4()}@example.com",
        password_hash="x", role=UserRole.ADMINISTRATOR, status=UserStatus.ACTIVE,
    )
    db.add(user)
    db.flush()
    return company, user


class _FakeUploadFile:
    """Minimal stand-in for fastapi.UploadFile -- storage.save_upload()
    only ever calls .filename / .content_type / (async) .read()."""

    def __init__(self, filename: str, content: bytes, content_type: str):
        self.filename = filename
        self.content_type = content_type
        self._content = content
        self._sent = False

    async def read(self, _size: int = -1) -> bytes:
        if self._sent:
            return b""
        self._sent = True
        return self._content

    async def seek(self, _pos: int) -> None:
        self._sent = False


def _real_pdf_bytes(lines: list[str]) -> bytes:
    """A genuinely valid, pypdf-readable single-page PDF with real text
    content -- not a fake '%PDF' byte string, since document_parser's
    extract_pdf_pages() needs real extractable text."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(612, 792))
    y = 750
    for line in lines:
        c.drawString(50, y, line)
        y -= 20
    c.showPage()
    c.save()
    return buf.getvalue()


def _real_xlsx_bytes(sheets: dict[str, list[list[str]]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 1 & 11. PDF-only tender still works / existing analysis path intact
# ---------------------------------------------------------------------------


def test_pdf_only_tender_analysis_unchanged(tmp_path, monkeypatch, db):
    monkeypatch.setattr(storage.settings, "storage_backend", "local")
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)

    company, user = _make_company_and_user(db)
    mission = Mission(
        id=uuid.uuid4(), company_id=company.id, user_id=user.id,
        mission_type="tender_evaluation", status=MissionStatus.CREATED,
    )
    db.add(mission)
    db.flush()

    pdf_bytes = _real_pdf_bytes(["Eligibility: Bidder must have completed similar works during the last 5 years."])
    doc_dir = tmp_path / str(company.id) / "documents"
    doc_dir.mkdir(parents=True)
    doc_path = doc_dir / f"{uuid.uuid4()}.pdf"
    doc_path.write_bytes(pdf_bytes)

    document = Document(
        id=uuid.uuid4(), company_id=company.id, uploaded_by=user.id, document_type="tender",
        file_name="tender.pdf", storage_path=str(doc_path.relative_to(tmp_path)),
    )
    db.add(document)
    db.flush()

    tender = Tender(
        id=uuid.uuid4(), mission_id=mission.id, tender_name="Single PDF Tender",
        organization="Test Org", uploaded_document=document.id,
        processing_status=DocumentProcessingStatus.PENDING.value,
    )
    db.add(tender)
    db.flush()
    # Backfilled by the real migration in production -- set explicitly here
    # since this test builds rows directly rather than via upload_tender().
    document.tender_id = tender.id
    document.document_role = "main"
    db.commit()

    result_tender, requirements = asyncio.run(
        tender_service.run_analysis(db, tender.id, company.id, provider="mock")
    )

    assert result_tender.processing_status == DocumentProcessingStatus.COMPLETED.value
    assert len(requirements) == 1
    req = requirements[0]
    assert req.requirement_type.value == "eligibility"
    # Backward-compat guarantee: a single-PDF tender's requirement carries a
    # real PDF page number, identical in shape to pre-multi-document results.
    assert req.source_page == 1
    assert req.source_document_id == document.id
    assert req.source_location is None


def test_upload_tender_links_document_via_new_general_relationship(tmp_path, monkeypatch, db):
    """upload_tender() must set the new Document.tender_id/document_role
    columns (not just the legacy Tender.uploaded_document) -- this is what
    lets run_analysis()'s general 'find every attached document' query work
    without any special-case fallback."""
    monkeypatch.setattr(storage.settings, "storage_backend", "local")
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)

    company, user = _make_company_and_user(db)
    pdf_bytes = _real_pdf_bytes(["Technical Requirement: ISO 9001 certification required."])
    file = _FakeUploadFile("tender.pdf", pdf_bytes, "application/pdf")

    _mission, tender = asyncio.run(
        tender_service.upload_tender(db, company.id, user.id, file, "T1", "Org", None, "IT Services")
    )

    document = db.get(Document, tender.uploaded_document)
    assert document.tender_id == tender.id
    assert document.document_role == "main"


# ---------------------------------------------------------------------------
# 2 & 3. XLS / XLSX accepted, 4. unsupported types still rejected
# ---------------------------------------------------------------------------


def test_xlsx_content_type_accepted():
    storage.validate_file_type(
        "tech.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )  # must not raise


def test_xls_content_type_accepted():
    storage.validate_file_type("tech.xls", "application/vnd.ms-excel")  # must not raise


def test_xls_octet_stream_leniency_accepted():
    """Real CPPP portals commonly serve .xls/.xlsx with a generic
    application/octet-stream content-type -- confirmed against the real
    tender.pdf/tech.xls/BOQ_969057.xls this feature was built against."""
    storage.validate_file_type("BOQ_969057.xls", "application/octet-stream")  # must not raise


def test_unsupported_extension_still_rejected():
    with pytest.raises(UnsupportedFileTypeError):
        storage.validate_file_type("malware.exe", "application/octet-stream")


def test_pdf_octet_stream_still_rejected():
    """The leniency is scoped to spreadsheets only -- a PDF with a wrong
    content-type is still rejected exactly as before."""
    with pytest.raises(UnsupportedFileTypeError):
        storage.validate_file_type("tender.pdf", "application/octet-stream")


# ---------------------------------------------------------------------------
# 6. Spreadsheet parsing: single sheet, multiple sheets, empty sheet skipped
# ---------------------------------------------------------------------------


def test_extract_spreadsheet_sheets_xlsx_multiple_sheets_and_empty_sheet(tmp_path):
    path = tmp_path / "tech.xlsx"
    path.write_bytes(
        _real_xlsx_bytes(
            {
                "Sheet1": [["Requirement", "Detail"], ["Eligibility", "5 years experience required."]],
                "EmptySheet": [],
                "Sheet2": [["Certification", "ISO 9001"]],
            }
        )
    )

    sheets = document_parser.extract_spreadsheet_sheets(path, ".xlsx")

    names = [name for name, _text in sheets]
    assert names == ["Sheet1", "Sheet2"]  # empty sheet dropped
    assert "Eligibility | 5 years experience required." in dict(sheets)["Sheet1"]
    assert "Certification | ISO 9001" in dict(sheets)["Sheet2"]


def test_extract_text_dispatches_xlsx_with_sheet_headers(tmp_path):
    path = tmp_path / "tech.xlsx"
    path.write_bytes(_real_xlsx_bytes({"Sheet1": [["Requirement", "Detail"], ["A", "B"]]}))

    parsed = document_parser.extract_text(path, ".xlsx")

    assert "=== Sheet1 ===" in parsed.text
    assert parsed.used_ocr is False


def test_unsupported_spreadsheet_extension_raises(tmp_path):
    with pytest.raises(ValueError):
        document_parser.extract_spreadsheet_sheets(tmp_path / "x.csv", ".csv")


# ---------------------------------------------------------------------------
# 7, 8, 9. Combined extraction, source traceability, financial exclusion
# ---------------------------------------------------------------------------


def test_analyze_tender_combines_pdf_and_spreadsheet_with_traceability(tmp_path):
    pdf_path = tmp_path / "tender.pdf"
    pdf_path.write_bytes(
        _real_pdf_bytes(["Eligibility: Bidder must have completed similar works during the last 5 years."])
    )

    xlsx_path = tmp_path / "tech.xlsx"
    xlsx_path.write_bytes(
        _real_xlsx_bytes({"Sheet1": [["Technical Requirement: ISO 9001 certification mandatory."]]})
    )

    pdf_doc_id, xlsx_doc_id = uuid.uuid4(), uuid.uuid4()
    sources = [
        TenderSourceDocument(document_id=pdf_doc_id, file_name="tender.pdf", document_role="main", file_path=pdf_path),
        TenderSourceDocument(document_id=xlsx_doc_id, file_name="tech.xlsx", document_role="technical", file_path=xlsx_path),
    ]

    results = asyncio.run(analyze_tender(sources, provider="mock"))

    by_type = {r.requirement_type: r for r in results}
    assert "eligibility" in by_type
    assert "technical" in by_type

    pdf_req = by_type["eligibility"]
    assert pdf_req.source_document_id == pdf_doc_id
    assert pdf_req.source_page == 1  # real PDF page number, not a raw unit index
    assert pdf_req.source_location is None

    xlsx_req = by_type["technical"]
    assert xlsx_req.source_document_id == xlsx_doc_id
    assert xlsx_req.source_page is None  # spreadsheets are not paginated
    assert xlsx_req.source_location == "Sheet: Sheet1"


def test_analyze_tender_excludes_financial_role_documents(tmp_path):
    pdf_path = tmp_path / "tender.pdf"
    pdf_path.write_bytes(_real_pdf_bytes(["Eligibility: Must be a registered contractor."]))

    boq_path = tmp_path / "BOQ_969057.xlsx"
    boq_path.write_bytes(
        _real_xlsx_bytes({"BOQ": [["Eligibility: This line must never reach the LLM (financial-role document)."]]})
    )

    sources = [
        TenderSourceDocument(document_id=uuid.uuid4(), file_name="tender.pdf", document_role="main", file_path=pdf_path),
        TenderSourceDocument(document_id=uuid.uuid4(), file_name="BOQ_969057.xlsx", document_role="financial", file_path=boq_path),
    ]

    results = asyncio.run(analyze_tender(sources, provider="mock"))

    descriptions = [r.description for r in results]
    assert any("registered contractor" in (d or "") for d in descriptions)
    assert not any("must never reach the LLM" in (d or "") for d in descriptions)


def test_analyze_tender_raises_when_no_sources():
    with pytest.raises(ValueError):
        asyncio.run(analyze_tender([], provider="mock"))


def test_analyze_tender_raises_when_only_financial_documents(tmp_path):
    boq_path = tmp_path / "BOQ.xlsx"
    boq_path.write_bytes(_real_xlsx_bytes({"BOQ": [["Item", "Amount"], ["Item 1", "100000"]]}))
    sources = [
        TenderSourceDocument(document_id=uuid.uuid4(), file_name="BOQ.xlsx", document_role="financial", file_path=boq_path),
    ]
    with pytest.raises(ValueError):
        asyncio.run(analyze_tender(sources, provider="mock"))


# ---------------------------------------------------------------------------
# 5. Multiple documents attached to one Tender / role inference
# ---------------------------------------------------------------------------


def test_add_tender_document_attaches_and_infers_role(tmp_path, monkeypatch, db):
    monkeypatch.setattr(storage.settings, "storage_backend", "local")
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)

    company, user = _make_company_and_user(db)
    main_pdf = _FakeUploadFile("tender.pdf", _real_pdf_bytes(["Eligibility: X"]), "application/pdf")
    _mission, tender = asyncio.run(
        tender_service.upload_tender(db, company.id, user.id, main_pdf, "T1", "Org", None, "IT Services")
    )

    tech_xls = _FakeUploadFile(
        "tech.xls", _real_xlsx_bytes({"S": [["a"]]}),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    boq_xls = _FakeUploadFile(
        "BOQ_969057.xls", _real_xlsx_bytes({"S": [["a"]]}),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    tech_doc = asyncio.run(tender_service.add_tender_document(db, tender.id, company.id, user.id, tech_xls))
    boq_doc = asyncio.run(tender_service.add_tender_document(db, tender.id, company.id, user.id, boq_xls))

    assert tech_doc.document_role == "technical"  # inferred from filename
    assert boq_doc.document_role == "financial"  # inferred from filename ("boq")

    documents = tender_service.list_tender_documents(db, tender.id, company.id)
    assert {d.file_name for d in documents} == {"tender.pdf", "tech.xls", "BOQ_969057.xls"}


def test_add_tender_document_respects_explicit_role_override(tmp_path, monkeypatch, db):
    monkeypatch.setattr(storage.settings, "storage_backend", "local")
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)

    company, user = _make_company_and_user(db)
    main_pdf = _FakeUploadFile("tender.pdf", _real_pdf_bytes(["Eligibility: X"]), "application/pdf")
    _mission, tender = asyncio.run(
        tender_service.upload_tender(db, company.id, user.id, main_pdf, "T1", "Org", None, "IT Services")
    )

    annexure = _FakeUploadFile("random_name.xlsx", _real_xlsx_bytes({"S": [["a"]]}), "application/octet-stream")
    doc = asyncio.run(
        tender_service.add_tender_document(db, tender.id, company.id, user.id, annexure, document_role="annexure")
    )
    assert doc.document_role == "annexure"


def test_add_tender_document_raises_not_found_for_wrong_company(tmp_path, monkeypatch, db):
    monkeypatch.setattr(storage.settings, "storage_backend", "local")
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)

    company_a, user_a = _make_company_and_user(db)
    company_b, _user_b = _make_company_and_user(db)
    main_pdf = _FakeUploadFile("tender.pdf", _real_pdf_bytes(["Eligibility: X"]), "application/pdf")
    _mission, tender = asyncio.run(
        tender_service.upload_tender(db, company_a.id, user_a.id, main_pdf, "T1", "Org", None, "IT Services")
    )

    other_doc = _FakeUploadFile("tech.xls", _real_xlsx_bytes({"S": [["a"]]}), "application/vnd.ms-excel")
    with pytest.raises(NotFoundError):
        asyncio.run(tender_service.add_tender_document(db, tender.id, company_b.id, user_a.id, other_doc))


# ---------------------------------------------------------------------------
# 10. Company isolation
# ---------------------------------------------------------------------------


def test_list_tender_documents_is_company_scoped(tmp_path, monkeypatch, db):
    monkeypatch.setattr(storage.settings, "storage_backend", "local")
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)

    company_a, user_a = _make_company_and_user(db)
    company_b, _user_b = _make_company_and_user(db)

    main_pdf = _FakeUploadFile("tender.pdf", _real_pdf_bytes(["Eligibility: X"]), "application/pdf")
    _mission, tender = asyncio.run(
        tender_service.upload_tender(db, company_a.id, user_a.id, main_pdf, "T1", "Org", None, "IT Services")
    )

    with pytest.raises(NotFoundError):
        tender_service.list_tender_documents(db, tender.id, company_b.id)


def test_delete_document_blocked_when_attached_via_new_relationship_only(tmp_path, monkeypatch, db):
    """document_service.delete_document()'s blocking check was generalized
    to cover Document.tender_id, not just the legacy Tender.uploaded_document
    -- an additional attached document (not the 'main' one) must also be
    protected from deletion while its tender is active."""
    monkeypatch.setattr(storage.settings, "storage_backend", "local")
    monkeypatch.setattr(storage, "STORAGE_ROOT", tmp_path)

    company, user = _make_company_and_user(db)
    main_pdf = _FakeUploadFile("tender.pdf", _real_pdf_bytes(["Eligibility: X"]), "application/pdf")
    _mission, tender = asyncio.run(
        tender_service.upload_tender(db, company.id, user.id, main_pdf, "T1", "Org", None, "IT Services")
    )
    tech_xls = _FakeUploadFile("tech.xls", _real_xlsx_bytes({"S": [["a"]]}), "application/vnd.ms-excel")
    tech_doc = asyncio.run(tender_service.add_tender_document(db, tender.id, company.id, user.id, tech_xls))

    # tech_doc is NOT Tender.uploaded_document (that's still the main PDF) --
    # only the new tender_id relationship protects it.
    assert tender.uploaded_document != tech_doc.id
    with pytest.raises(ConflictError):
        document_service.delete_document(db, tech_doc.id, company.id)
